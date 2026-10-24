"""
Helper to create native .pdf and .docx sample brief files in data/sample_briefs/.
Allows evaluators and testers to test file drag-and-drop with genuine binary document formats.
"""

from pathlib import Path
import docx
import fitz  # PyMuPDF

def create_sample_docx():
    txt_path = Path("data/sample_briefs/ecommerce_omnichannel.txt")
    if not txt_path.exists():
        return
    text = txt_path.read_text(encoding="utf-8")

    doc = docx.Document()
    doc.add_heading("PROJECT BRIEF: Modern Omnichannel E-Commerce Platform", level=0)
    for line in text.split("\n"):
        line = line.strip()
        if not line:
            continue
        if line.endswith(":") and len(line) < 40:
            doc.add_heading(line, level=1)
        elif line.startswith("- ") or line.startswith("• "):
            doc.add_paragraph(line[2:], style='List Bullet')
        elif line.startswith("1.") or line.startswith("2.") or line.startswith("3."):
            doc.add_paragraph(line, style='List Number')
        else:
            doc.add_paragraph(line)

    out_path = Path("data/sample_briefs/ecommerce_brief.docx")
    doc.save(out_path)
    print(f"Created sample DOCX: {out_path}")

def create_sample_pdf():
    txt_path = Path("data/sample_briefs/healthcare_telemed.txt")
    if not txt_path.exists():
        return
    text = txt_path.read_text(encoding="utf-8")

    pdf_doc = fitz.open()
    page = pdf_doc.new_page(width=595, height=842) # A4
    rect = fitz.Rect(50, 50, 545, 792)
    
    formatted_text = f"PROJECT BRIEF: CareConnect Telehealth Platform\n\n{text}"
    page.insert_textbox(rect, formatted_text, fontsize=11, fontname="helv", color=(0.1, 0.1, 0.1))
    
    out_path = Path("data/sample_briefs/healthcare_telemed_brief.pdf")
    pdf_doc.save(out_path)
    pdf_doc.close()
    print(f"Created sample PDF: {out_path}")

if __name__ == "__main__":
    create_sample_docx()
    create_sample_pdf()
