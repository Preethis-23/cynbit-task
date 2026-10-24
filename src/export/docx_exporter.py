"""
Proposal Exporter for Microsoft Word (.docx) and Markdown (.md).
Generates client-ready executive proposals with styled headers, structured tables, and disclaimers.
"""

import io
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from src.schemas.proposal import ProposalContent
from src.schemas.requirements import RequirementAnalysis
from src.schemas.estimation import CommercialAssessment
from src.schemas.risk import RiskAnalysisResult


class ProposalExporter:
    """Exports structured proposals to DOCX and Markdown."""

    @staticmethod
    def to_markdown(
        proposal: ProposalContent,
        analysis: RequirementAnalysis,
        commercial: CommercialAssessment,
        risk: RiskAnalysisResult,
    ) -> str:
        """Formats proposal as a clean, complete Markdown document."""
        md_lines = [
            f"# {proposal.title}",
            "",
            f"**Client:** {proposal.client_name}  ",
            f"**Prepared By:** {proposal.prepared_by}  ",
            f"**Date:** {proposal.date_created}  ",
            f"**Overall Complexity:** {risk.complexity_level} | **Risk Level:** {risk.risk_level}  ",
            "",
            "---",
            "",
            "## 1. Executive Summary",
            "",
            proposal.executive_summary,
            "",
            "## 2. Business Objectives & Problem Context",
            "",
            proposal.business_context_and_goals,
            "",
            "## 3. Scope of Work (Functional Deliverables)",
            "",
        ]

        for item in proposal.scope_of_work:
            md_lines.append(f"- {item}")

        md_lines.extend([
            "",
            "## 4. Out-of-Scope Exclusions",
            "",
        ])
        for exc in proposal.out_of_scope_exclusions:
            md_lines.append(f"- {exc}")

        md_lines.extend([
            "",
            "## 5. Phase-by-Phase Delivery Roadmap",
            "",
            "| Phase | Milestone | Deliverables | Est. Weeks |",
            "| :--- | :--- | :--- | :--- |",
        ])
        for p in proposal.deliverables_by_phase:
            deliv_str = "<br>• ".join(p.deliverables)
            md_lines.append(f"| **{p.phase_name}** | {p.milestone} | • {deliv_str} | {p.estimated_duration_weeks:.1f} wks |")

        md_lines.extend([
            "",
            "## 6. Recommended Technology Architecture",
            "",
            proposal.recommended_tech_stack_summary,
            "",
            "### Architectural Decisions:",
        ])
        for dec in proposal.key_architectural_decisions:
            md_lines.append(f"- {dec}")

        md_lines.extend([
            "",
            "## 7. Timeline & Resource Structure",
            "",
            f"**Timeline Summary:** {proposal.timeline_summary}",
            "",
            f"**Team Structure:** {proposal.recommended_team_structure}",
            "",
            "## 8. Commercial & Pricing Estimate",
            "",
            f"**Indicative Quotation:** ${commercial.quotation_min:,.0f} - ${commercial.quotation_max:,.0f} "
            f"(Target Estimate: **${commercial.quotation_likely:,.0f}**)",
            "",
            f"**Projected Internal Cost:** ${commercial.internal_cost_likely:,.0f}  ",
            f"**Target Profit Margin:** ~{commercial.profit_margin_likely_pct:.1f}%  ",
            "",
            f"**Payment Structure:** {proposal.pricing_options_summary}",
            "",
            "## 9. Key Project Assumptions",
            "",
        ])
        for asm in proposal.assumptions_and_prerequisites:
            md_lines.append(f"- {asm}")

        md_lines.extend([
            "",
            "## 10. Client Clarification Questions",
            "",
        ])
        for q in proposal.client_clarification_questions:
            md_lines.append(f"1. {q}")

        md_lines.extend([
            "",
            "## 11. Next Steps",
            "",
        ])
        for step in proposal.next_steps:
            md_lines.append(f"- {step}")

        md_lines.extend([
            "",
            "---",
            "",
            "> **Disclaimer & Human Review:**  ",
            f"> {proposal.expert_review_disclaimer}",
        ])

        return "\n".join(md_lines)

    @classmethod
    def to_docx(
        cls,
        proposal: ProposalContent,
        analysis: RequirementAnalysis,
        commercial: CommercialAssessment,
        risk: RiskAnalysisResult,
    ) -> io.BytesIO:
        """Builds a formatted Word document."""
        doc = Document()

        # Styles
        # Normal text font
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Calibri'
        font.size = Pt(11)
        font.color.rgb = RGBColor(0x33, 0x33, 0x33)

        # Header Title
        title_p = doc.add_paragraph()
        title_run = title_p.add_run(proposal.title)
        title_run.font.size = Pt(22)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(0x0f, 0x2b, 0x46)  # Deep Navy
        title_p.paragraph_format.space_after = Pt(4)

        # Subtitle metadata table
        meta_table = doc.add_table(rows=2, cols=2)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_table.rows[0].cells[0].text = f"Client: {proposal.client_name}"
        meta_table.rows[0].cells[1].text = f"Date: {proposal.date_created}"
        meta_table.rows[1].cells[0].text = f"Prepared By: {proposal.prepared_by}"
        meta_table.rows[1].cells[1].text = f"Complexity: {risk.complexity_level} | Risk: {risk.risk_level}"

        doc.add_paragraph().paragraph_format.space_after = Pt(12)

        # 1. Executive Summary
        h1 = doc.add_heading("1. Executive Summary", level=1)
        h1.paragraph_format.space_before = Pt(14)
        doc.add_paragraph(proposal.executive_summary)

        # 2. Business Objective
        doc.add_heading("2. Business Objectives & Context", level=1)
        doc.add_paragraph(proposal.business_context_and_goals)

        # 3. Scope of Work
        doc.add_heading("3. Scope of Work", level=1)
        for item in proposal.scope_of_work:
            doc.add_paragraph(item, style='List Bullet')

        # 4. Out of Scope
        doc.add_heading("4. Exclusions & Out of Scope", level=1)
        for exc in proposal.out_of_scope_exclusions:
            doc.add_paragraph(exc, style='List Bullet')

        # 5. Phased Delivery Roadmap Table
        doc.add_heading("5. Delivery Roadmap & Milestones", level=1)
        table = doc.add_table(rows=1, cols=4)
        table.style = 'Light Shading Accent 1'
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Phase"
        hdr_cells[1].text = "Milestone"
        hdr_cells[2].text = "Key Deliverables"
        hdr_cells[3].text = "Duration"

        for p in proposal.deliverables_by_phase:
            row_cells = table.add_row().cells
            row_cells[0].text = p.phase_name
            row_cells[1].text = p.milestone
            row_cells[2].text = "\n• ".join([""] + p.deliverables).strip()
            row_cells[3].text = f"{p.estimated_duration_weeks:.0f} weeks"

        doc.add_paragraph().paragraph_format.space_after = Pt(10)

        # 6. Technology Architecture
        doc.add_heading("6. Recommended Technology Architecture", level=1)
        doc.add_paragraph(proposal.recommended_tech_stack_summary)
        for dec in proposal.key_architectural_decisions:
            doc.add_paragraph(dec, style='List Bullet')

        # 7. Timeline & Team
        doc.add_heading("7. Timeline & Team Structure", level=1)
        doc.add_paragraph(f"Timeline Summary: {proposal.timeline_summary}")
        doc.add_paragraph(f"Recommended Team: {proposal.recommended_team_structure}")

        # 8. Commercial Assessment Table
        doc.add_heading("8. Commercial Assessment & Pricing", level=1)
        comm_table = doc.add_table(rows=4, cols=2)
        comm_table.style = 'Table Grid'
        comm_table.rows[0].cells[0].text = "Quotation Range (Indicative)"
        comm_table.rows[0].cells[1].text = f"${commercial.quotation_min:,.0f} - ${commercial.quotation_max:,.0f} (Target: ${commercial.quotation_likely:,.0f})"
        comm_table.rows[1].cells[0].text = "Estimated Engineering Effort"
        comm_table.rows[1].cells[1].text = f"{commercial.internal_cost_likely / commercial.hourly_internal_cost:,.0f} Hours"
        comm_table.rows[2].cells[0].text = "Expected Profit Margin"
        comm_table.rows[2].cells[1].text = f"~{commercial.profit_margin_likely_pct:.1f}%"
        comm_table.rows[3].cells[0].text = "Billing Model"
        comm_table.rows[3].cells[1].text = proposal.pricing_options_summary

        doc.add_paragraph().paragraph_format.space_after = Pt(10)

        # 9. Key Assumptions
        doc.add_heading("9. Key Assumptions", level=1)
        for asm in proposal.assumptions_and_prerequisites:
            doc.add_paragraph(asm, style='List Bullet')

        # 10. Clarification Questions
        doc.add_heading("10. Clarification Questions for Client", level=1)
        for q in proposal.client_clarification_questions:
            doc.add_paragraph(q, style='List Number')

        # 11. Next Steps
        doc.add_heading("11. Next Steps", level=1)
        for step in proposal.next_steps:
            doc.add_paragraph(step, style='List Bullet')

        # Disclaimer
        doc.add_paragraph().paragraph_format.space_after = Pt(16)
        disc_p = doc.add_paragraph()
        disc_run = disc_p.add_run(f"DISCLAIMER: {proposal.expert_review_disclaimer}")
        disc_run.font.italic = True
        disc_run.font.size = Pt(9)
        disc_run.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

        output = io.BytesIO()
        doc.save(output)
        output.seek(0)
        return output
