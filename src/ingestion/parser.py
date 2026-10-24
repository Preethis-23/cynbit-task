"""
Robust document parser for PDF, DOCX, and Plain Text briefs.
Normalizes whitespace, validates safety, and generates extraction metadata.
"""

import io
import re
from pathlib import Path
from typing import Union, Optional

from src.schemas.document import DocumentMetadata, IngestedDocument
from config.settings import settings


class IngestionError(Exception):
    """Raised when document ingestion or parsing fails."""
    pass


class DocumentParser:
    """Handles parsing and sanitization of client requirement documents."""

    def __init__(self, max_file_size_mb: int = settings.max_file_size_mb):
        self.max_size_bytes = max_file_size_mb * 1024 * 1024

    @staticmethod
    def normalize_text(text: str) -> str:
        """
        Cleans extracted text:
        - Normalizes line breaks
        - Removes control characters except newlines/tabs
        - Collapses 3+ consecutive newlines into 2
        - Strips trailing/leading whitespaces
        """
        if not text:
            return ""

        # Replace carriage returns
        text = text.replace("\r\n", "\n").replace("\r", "\n")
        # Remove non-printable control characters except \t and \n
        text = "".join(ch for ch in text if ch == "\n" or ch == "\t" or ch.isprintable())
        # Replace non-breaking spaces
        text = text.replace("\u00a0", " ")
        # Collapse multiple empty lines
        text = re.sub(r"\n{3,}", "\n\n", text)
        # Strip each line
        lines = [line.strip() for line in text.split("\n")]
        return "\n".join(lines).strip()

    def parse_plain_text(self, text: str, source_type: str = "text", filename: Optional[str] = None) -> IngestedDocument:
        """Parse raw text string directly."""
        warnings: list[str] = []
        if not text or not text.strip():
            raise IngestionError("The provided requirement text is empty. Please enter or upload a valid brief.")

        byte_length = len(text.encode("utf-8"))
        if byte_length > self.max_size_bytes:
            raise IngestionError(f"Text content exceeds the {settings.max_file_size_mb} MB limit.")

        cleaned = self.normalize_text(text)
        words = len(cleaned.split())
        chars = len(cleaned)
        paragraphs = len([p for p in cleaned.split("\n\n") if p.strip()])

        if words < 15:
            warnings.append("The document text is very short (<15 words). High risk of missing requirement context.")

        metadata = DocumentMetadata(
            source_type=source_type if source_type in ["text", "pdf", "docx", "sample"] else "text",
            filename=filename,
            char_count=chars,
            word_count=words,
            paragraph_count=paragraphs,
            warnings=warnings,
        )

        return IngestedDocument(
            raw_text=text,
            cleaned_text=cleaned,
            metadata=metadata,
        )

    def parse_pdf(self, file_content: Union[bytes, io.BytesIO], filename: str = "document.pdf") -> IngestedDocument:
        """Extract text from a PDF file using PyMuPDF (fitz) with PyPDF fallback."""
        if isinstance(file_content, io.BytesIO):
            data = file_content.getvalue()
        else:
            data = file_content

        if len(data) == 0:
            raise IngestionError("Uploaded PDF file is empty.")
        if len(data) > self.max_size_bytes:
            raise IngestionError(f"PDF file exceeds the {settings.max_file_size_mb} MB limit.")

        extracted_text = ""
        warnings: list[str] = []

        # Attempt 1: PyMuPDF (fast, high fidelity)
        try:
            import fitz
            doc = fitz.open(stream=data, filetype="pdf")
            pages_text = []
            for page_num in range(len(doc)):
                page = doc.load_page(page_num)
                pages_text.append(page.get_text())
            doc.close()
            extracted_text = "\n\n".join(pages_text)
        except Exception as fitz_err:
            # Attempt 2: PyPDF fallback
            try:
                import pypdf
                reader = pypdf.PdfReader(io.BytesIO(data))
                pages_text = []
                for page in reader.pages:
                    text = page.extract_text()
                    if text:
                        pages_text.append(text)
                extracted_text = "\n\n".join(pages_text)
            except Exception as pypdf_err:
                raise IngestionError(f"Failed to extract text from PDF: {fitz_err}")

        cleaned = self.normalize_text(extracted_text)
        if not cleaned:
            raise IngestionError(
                "No readable text found in the PDF. It might be scanned/image-only or password protected."
            )

        words = len(cleaned.split())
        if words < 20:
            warnings.append("Low word count detected in PDF. Some pages may contain images or tables that were skipped.")

        metadata = DocumentMetadata(
            source_type="pdf",
            filename=filename,
            char_count=len(cleaned),
            word_count=words,
            paragraph_count=len([p for p in cleaned.split("\n\n") if p.strip()]),
            warnings=warnings,
        )

        return IngestedDocument(
            raw_text=extracted_text,
            cleaned_text=cleaned,
            metadata=metadata,
        )

    def parse_docx(self, file_content: Union[bytes, io.BytesIO], filename: str = "document.docx") -> IngestedDocument:
        """Extract text from a DOCX file using python-docx."""
        if isinstance(file_content, io.BytesIO):
            data = file_content.getvalue()
        else:
            data = file_content

        if len(data) == 0:
            raise IngestionError("Uploaded DOCX file is empty.")
        if len(data) > self.max_size_bytes:
            raise IngestionError(f"DOCX file exceeds the {settings.max_file_size_mb} MB limit.")

        try:
            import docx
            doc = docx.Document(io.BytesIO(data))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        paragraphs.append(" | ".join(row_cells))

            extracted_text = "\n\n".join(paragraphs)
        except Exception as err:
            raise IngestionError(f"Failed to read DOCX file: {err}")

        cleaned = self.normalize_text(extracted_text)
        if not cleaned:
            raise IngestionError("DOCX file does not contain readable text.")

        metadata = DocumentMetadata(
            source_type="docx",
            filename=filename,
            char_count=len(cleaned),
            word_count=len(cleaned.split()),
            paragraph_count=len([p for p in cleaned.split("\n\n") if p.strip()]),
            warnings=[],
        )

        return IngestedDocument(
            raw_text=extracted_text,
            cleaned_text=cleaned,
            metadata=metadata,
        )

    def parse_file(self, file_content: Union[bytes, io.BytesIO], filename: str) -> IngestedDocument:
        """Dispatch parsing based on file extension."""
        ext = Path(filename).suffix.lower().lstrip(".")
        if ext == "pdf":
            return self.parse_pdf(file_content, filename=filename)
        elif ext in ["docx", "doc"]:
            return self.parse_docx(file_content, filename=filename)
        elif ext in ["txt", "md", "text"]:
            if isinstance(file_content, io.BytesIO):
                raw_bytes = file_content.getvalue()
            else:
                raw_bytes = file_content
            try:
                text = raw_bytes.decode("utf-8")
            except UnicodeDecodeError:
                text = raw_bytes.decode("latin-1", errors="replace")
            return self.parse_plain_text(text, source_type="text", filename=filename)
        else:
            raise IngestionError(
                f"Unsupported file format '.{ext}'. Supported formats are: {', '.join(settings.allowed_extensions)}"
            )
